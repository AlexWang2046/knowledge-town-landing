from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageOps
import requests, os, tempfile

OUT='docs/KnowledgeTown_OPC_BP_15slides_latest.pptx'
W,H=13.333,7.5
CREAM='F7F3E8'; GREEN='214C3C'; GREEN2='315F4C'; SAGE='DDE7D3'; SAGE2='E9EFE2'; GOLD='D7A63B'; ORANGE='DD8355'; INK='20332B'; MUTED='6F786F'; WHITE='FFFDF8'; BORDER='D9D4C6'; BLUE='6D8E9B'
FONT='Noto Sans CJK SC'
RAW='https://raw.githubusercontent.com/Abigale365/Knowledge_town/main/dist/assets/'
ASSETS={
 'grove':RAW+'world-grove.webp','academy':RAW+'academy-map.webp','archive':RAW+'archive.webp',
 'seaside':RAW+'away-seaside.png','winter':RAW+'away-winter.png','workshop':RAW+'workshop.webp',
 'ib_site':'https://media.canva.com/v2/image-resize/format:JPG/height:200/quality:75/uri:ifs%3A%2F%2FM%2F20717ad1-b8a5-4faf-99f9-bbb35ea90a1b/watermark:F/width:117?csig=AAAAAAAAAAAAAAAAAAAAAEKPNsztMGIwAe2qdT_93bYNQt4EmC1n4kuh1xXpwIWK&exp=1792868007&osig=AAAAAAAAAAAAAAAAAAAAAM_YpztKjS5IPxHq3o1qle3vmrQcz4H0Jc3iX27ZAglk&signer=media-rpc&token=AAIAAU0AJDIwNzE3YWQxLWI4YTUtNGZhZi05OWY5LWJiYjM1ZWE5MGExYgAAAAABoW9CvFhvD7CqEqOdaRBEDRgBwi_Camvjug4jVrKBGg5h8y0HBw&x-canva-quality=thumbnail',
 'ib_xhs':'https://media.canva.com/v2/image-resize/format:JPG/height:200/quality:75/uri:ifs%3A%2F%2FM%2Fa005027b-8c19-44b3-bfbb-b1153da8fb0d/watermark:F/width:106?csig=AAAAAAAAAAAAAAAAAAAAANwkYYrVZwHXyc6N0-k_5ShNz5jGi3HgHPa2GXjhiaFB&exp=1792868591&osig=AAAAAAAAAAAAAAAAAAAAAJhD4enj2H47SQLNl-rsiXom3oAICwEMaRxrBe0-dsbt&signer=media-rpc&token=AAIAAU0AJGEwMDUwMjdiLThjMTktNDRiMy1iZmJiLWIxMTUzZGE4ZmIwZAAAAAABoW9LpZhYzOv9iqQZRbItzPuGuCfjYWrNxTp4LRN0fZc1-SNZRA&x-canva-quality=thumbnail'
}

def C(h): return RGBColor.from_string(h)
def bg(s,c=CREAM): s.background.fill.solid(); s.background.fill.fore_color.rgb=C(c)
def txt(s,t,x,y,w,h,z=20,c=INK,b=False,a=PP_ALIGN.LEFT,v=MSO_ANCHOR.TOP):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); f=sh.text_frame; f.clear(); f.word_wrap=True; f.margin_left=f.margin_right=Inches(.02); f.margin_top=f.margin_bottom=Inches(.02); f.vertical_anchor=v
    p=f.paragraphs[0]; p.text=t; p.alignment=a; r=p.runs[0]; r.font.name=FONT; r.font.size=Pt(z); r.font.bold=b; r.font.color.rgb=C(c); return sh
def title(s,t,sub='',k=''):
    if k: txt(s,k,.7,.3,3,.25,10,GREEN2,True)
    txt(s,t,.7,.65,11.4,.7,30,GREEN,True)
    if sub: txt(s,sub,.72,1.35,11,.42,13,MUTED)
    txt(s,'拾知小镇 · KNOWLEDGE TOWN',10.55,.3,2,.2,8,GREEN2,True,PP_ALIGN.RIGHT)
def card(s,x,y,w,h,hd,body='',fill=WHITE,accent=GREEN):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=C(fill); q.line.color.rgb=C(BORDER)
    if accent:
        a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x+.12),Inches(y+.18),Inches(.07),Inches(h-.36)); a.fill.solid(); a.fill.fore_color.rgb=C(accent); a.line.fill.background()
    txt(s,hd,x+.28,y+.18,w-.45,.34,17,GREEN,True)
    if body: txt(s,body,x+.28,y+.62,w-.48,h-.75,12,MUTED)
def pill(s,t,x,y,w,fill=SAGE2,c=GREEN,z=11):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.48)); q.fill.solid(); q.fill.fore_color.rgb=C(fill); q.line.fill.background(); txt(s,t,x,y+.03,w,.36,z,c,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
def stat(s,v,l,x,y,w=2.2,fill=SAGE2):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(1.15)); q.fill.solid(); q.fill.fore_color.rgb=C(fill); q.line.fill.background(); txt(s,v,x,y+.08,w,.5,27,GREEN,True,PP_ALIGN.CENTER); txt(s,l,x,y+.68,w,.28,10,MUTED,False,PP_ALIGN.CENTER)
def foot(s,n): txt(s,str(n),12.55,7.12,.25,.15,8,GREEN2,False,PP_ALIGN.RIGHT)
def dl(key):
    os.makedirs('/tmp/ktassets',exist_ok=True); ext='.png' if key in ('seaside','winter') else '.jpg'; p=f'/tmp/ktassets/{key}{ext}'
    if os.path.exists(p): return p
    try:
        r=requests.get(ASSETS[key],timeout=35); r.raise_for_status(); open(p,'wb').write(r.content); Image.open(p).verify(); return p
    except Exception as e: print('asset failed',key,e); return None
def pic(s,key,x,y,w,h):
    p=dl(key)
    if not p: return
    im=Image.open(p).convert('RGB'); target=(int(w*180),int(h*180)); im=ImageOps.fit(im,target,method=Image.Resampling.LANCZOS); q=f'/tmp/ktassets/{key}_fit_{int(x*10)}_{int(y*10)}.jpg'; im.save(q,quality=90); s.shapes.add_picture(q,Inches(x),Inches(y),Inches(w),Inches(h))

prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H); blank=prs.slide_layouts[6]
# 1
s=prs.slides.add_slide(blank); bg(s); pic(s,'grove',8.55,0,4.78,7.5); txt(s,'WaytoAGI 1024 · OPC赛道',.8,.75,4,.3,12,GREEN2,True); txt(s,'拾知小镇',.8,1.35,4.5,.6,36,GREEN,True); txt(s,'Knowledge Town',.8,1.95,4.5,.35,20,GREEN2); txt(s,'上传你的知识，生成一个\n可以学习的世界。',.8,3.0,7.2,1.15,31,GREEN,True); txt(s,'AI 将教材、讲义和知识资料，转化为可探索、可互动、可练习的游戏化学习体验。',.8,4.75,6.5,.8,15,MUTED); pill(s,'ONE PERSON + AI',.8,6.2,2,fill=GREEN,c=WHITE); foot(s,1)
# 2
s=prs.slides.add_slide(blank); bg(s); title(s,'有知识，不等于有好的学习体验','问题不是缺知识，而是知识很难被转化成好的学习体验。','WHY');
for i,(a,b) in enumerate([('教材','内容丰富，却难互动'),('PPT','信息完整，却难练习'),('讲义','适合阅读，难形成反馈'),('文档','知识分散，难形成学习闭环')]): card(s,.75+i*3.12,2.35,2.62,2.65,a,b,SAGE2 if i%2==0 else WHITE,[GREEN,ORANGE,GOLD,BLUE][i]);
pill(s,'静态资料 ≠ 互动学习产品',.8,6.25,4.2,fill=GREEN,c=WHITE,z=13); foot(s,2)
# 3
s=prs.slides.add_slide(blank); bg(s); title(s,'拾知小镇是什么','内容属于用户，平台负责把知识变成体验。','PRODUCT');
for i,(a,b) in enumerate([('用户自己的知识','教材 · 讲义 · 文章 · 培训资料'),('AI 加工','理解知识 · 设计学习 · 生成互动'),('游戏化学习世界','探索 · 练习 · 反馈 · 复习')]): card(s,.8+i*4.05,2.15,3.55,2.0,a,b,SAGE2 if i!=1 else WHITE,[GREEN,GOLD,ORANGE][i]);
pic(s,'academy',7.75,4.35,4.8,2.2); txt(s,'用户资料  →  AI理解与编排  →  互动学习世界',.9,5.0,6.3,.6,20,GREEN,True); foot(s,3)
# 4
s=prs.slides.add_slide(blank); bg(s); title(s,'别人生成“内容”，我们生成“体验”','不是摘要 / PPT / 题库生成器。','DIFFERENCE');
for i,(a,b) in enumerate([('地图','进入学习世界'),('场景','知识进入环境'),('任务','带着目标学习'),('互动','真正操作知识'),('反馈','练习 · 错题 · 复习')]): card(s,.72+i*2.5,2.45,2.08,2.25,a,b,SAGE2 if i%2==0 else WHITE,[GREEN,ORANGE,GOLD,BLUE,GREEN2][i]);
txt(s,'知识不只是被“看见”，而是被“使用”。',1.3,5.65,10.7,.6,24,GREEN,True,PP_ALIGN.CENTER); foot(s,4)
# 5
s=prs.slides.add_slide(blank); bg(s); title(s,'一份资料，怎么进入拾知小镇','生成结果先预览、核对，再进入课程。','FLOW');
for i,t in enumerate(['上传 / 粘贴资料','建立知识库','AI 生成初稿','用户预览与调整','选择互动方式','进入游戏化课程']): card(s,.65+i*2.08,2.2,1.82,1.45,str(i+1),t,SAGE2 if i%2==0 else WHITE,[GREEN,GOLD,ORANGE,BLUE,GREEN2,GOLD][i]);
pic(s,'archive',.9,4.25,4.7,2.1); card(s,5.95,4.25,6.1,2.1,'真实工作流','上传资料 → AI 整理 → 生成课程 → 学习回流\n\nAI 做重复工作，人保留内容与教学判断。',WHITE,GREEN); foot(s,5)
# 6
s=prs.slides.add_slide(blank); bg(s); title(s,'《AI 商业入门》已经做成一个学习世界','不是概念演示，而是一套可以直接体验的完整样例。','SAMPLE'); pic(s,'grove',7.5,1.8,5.0,4.85); card(s,.8,1.95,5.9,3.1,'真实样例｜AI 商业入门','用户输入自己的资料后，系统把知识组织成章节、互动与练习。\n\n它是拾知小镇平台能力的第一个完整样例。',WHITE,GREEN);
for i,(v,l) in enumerate([('6','章地图'),('139','知识单元'),('23','小游戏单元'),('11','互动类型')]): stat(s,v,l,.85+i*1.55,5.4,1.35); foot(s,6)
# 7
s=prs.slides.add_slide(blank); bg(s); title(s,'游戏不是套皮，底层是一套学习闭环','让知识经历理解、操作、应用和反馈。','LEARNING LOOP');
for i,(v,l) in enumerate([('理解','基础 + 示范'),('操作','互动练习'),('应用','独立应用'),('反馈','错题 + 复习')]): stat(s,v,l,.8+i*3.05,2.05,2.55,SAGE2 if i%2==0 else WHITE)
for i,t in enumerate(['基础','示范','引导练习','独立应用','正式练习','错题修复','章末测验']): pill(s,t,.58+i*1.78,4.45,1.52,WHITE if i%2 else SAGE2)
card(s,1.0,5.55,11.3,1.0,'完整学习闭环','基础学习 → 示范 → 引导练习 → 独立应用 → 正式练习 → 错题修复 → 复习 / 测验',GREEN,None); txt(s,'完整学习闭环',1.25,5.7,2.2,.25,14,WHITE,True); txt(s,'基础学习 → 示范 → 引导练习 → 独立应用 → 正式练习 → 错题修复 → 复习 / 测验',1.25,6.05,10.4,.25,12,WHITE); foot(s,7)
# 8
s=prs.slides.add_slide(blank); bg(s); title(s,'卖点一｜内容属于用户','平台不规定你学什么。','SELLING POINT 01');
for i,(a,b) in enumerate([('独立老师','教材 / 讲义 / 自编课程'),('知识创作者','文章 / 专栏 / 系列知识'),('培训师','培训材料 / 内部课程'),('企业培训','SOP / 产品知识 / 员工培训')]): card(s,.85+i*3.05,2.2,2.6,2.6,a,b,SAGE2 if i%2==0 else WHITE,[GREEN,GOLD,ORANGE,BLUE][i]);
txt(s,'不同人的知识，进入同一套学习生产系统。',1.0,5.65,11.2,.55,23,GREEN,True,PP_ALIGN.CENTER); foot(s,8)
# 9
s=prs.slides.add_slide(blank); bg(s); title(s,'卖点二｜不同知识，用不同方式学习','知识是什么，决定它应该怎么被学习。','SELLING POINT 02'); card(s,.8,2,3.6,3.9,'知识类型','概念\n规则\n流程\n数据\n判断\n参数',SAGE2,GREEN); card(s,4.85,2,3.6,3.9,'适合的学习方式','分类 / 匹配\n排序 / 拖拽\n情境判断\n参数操作\n规则实验',WHITE,GOLD); card(s,8.9,2,3.6,3.9,'用户可以选择 / 组合','AI 可以推荐\n用户可以修改\n多种互动可组合\n最终由人决定',SAGE2,ORANGE); foot(s,9)
# 10
s=prs.slides.add_slide(blank); bg(s); title(s,'卖点三｜AI 做重活，人保留决定权','不是“一键生成以后就不能改”。','SELLING POINT 03'); card(s,.9,2,5.25,4.45,'AI 负责','理解资料\n整理知识\n生成初稿\n推荐学习方式',SAGE2,GREEN); card(s,7.15,2,5.25,4.45,'用户负责','判断内容\n选择互动\n修改结构\n审核结果',WHITE,GOLD); txt(s,'AI 提高生产效率，人保留教学判断。',3.1,6.55,7.1,.4,22,GREEN,True,PP_ALIGN.CENTER); foot(s,10)
# 11
s=prs.slides.add_slide(blank); bg(s); title(s,'卖点四｜同一套知识，可以进入不同学习世界','选现成模板 · 自己制作 · 上传保存 · 反复复用','SELLING POINT 04');
for i,(k,l) in enumerate([('grove','森林模板'),('seaside','海边模板'),('winter','冬日模板'),('workshop','自制模板示例')]): pic(s,k,.8+i*3.1,2.0,2.7,3.2); pill(s,l,.95+i*3.1,5.45,2.4,WHITE if i%2 else SAGE2); foot(s,11)
# 12
s=prs.slides.add_slide(blank); bg(s); title(s,'第一批用户：已经有知识，但缺产品能力的人','他们缺的不是内容，而是一套把内容变成学习产品的能力。','WHO');
for i,(a,b) in enumerate([('独立老师','有课程，没有开发团队'),('知识创作者','有内容，希望产品化'),('培训师 / 小型机构','有资料，需要更好的学习体验'),('企业培训','SOP / 产品知识 / 员工学习')]): card(s,.9+(i%2)*6.15,2.05+(i//2)*2.05,5.55,1.65,a,b,SAGE2 if i in (0,3) else WHITE,[GREEN,GOLD,ORANGE,BLUE][i]); foot(s,12)
# 13
s=prs.slides.add_slide(blank); bg(s); title(s,'为什么一个人，也能完成更多环节','一个人不是完成所有工作，而是拥有过去一支小团队才具备的生产能力。','OPC'); card(s,.85,2,5.35,3.65,'过去｜多人协作','内容整理\n教研\n交互设计\n视觉设计\n开发\n测试',WHITE,ORANGE); card(s,7.15,2,5.35,3.65,'现在｜人 + AI + 拾知小镇','人负责：需求判断 / 产品方向 / 教学判断 / 最终审核\n\nAI + 系统负责：整理 / 生成 / 编排 / 互动 / 重复生产',SAGE2,GREEN); pill(s,'ONE PERSON = SMALL TEAM CAPABILITY',3.7,6.15,5.9,GREEN,WHITE,13); foot(s,13)
# 14
s=prs.slides.add_slide(blank); bg(s); title(s,'已验证案例｜IB Practice Lab','团队成员已跑通垂直学习网站：从小红书内容获客，到网站学习与付费。','COMMERCIAL VALIDATION'); pic(s,'ib_site',.8,2,3.25,3.55); pic(s,'ib_xhs',4.3,2,3.25,3.55); stat(s,'3 个月','上线验证',8.0,2.15,2.0); stat(s,'1 万+','营业额',10.3,2.15,2.0,WHITE); card(s,8.0,3.75,4.3,1.55,'IB Practice Lab','IB 全科历年真题题库\n在线答题 · 自动判分 · 进度追踪 · 模考训练',SAGE2,GREEN); card(s,8.0,5.45,4.3,.95,'真实获客路径','小红书内容获客 → 网站承接 → 用户付费',GREEN,None); txt(s,'小红书主打：告别翻 PDF，直接组卷练习。',.9,5.85,6.6,.3,13,GREEN,True); txt(s,'团队成员既有项目，并非拾知小镇营收。',.9,6.4,6.7,.25,10,MUTED); foot(s,14)
# 15
s=prs.slides.add_slide(blank); bg(s); pic(s,'grove',8.8,0,4.53,7.5); txt(s,'上传你的知识，\n生成一个可以学习的世界。',.85,1.2,7,1.5,33,GREEN,True); txt(s,'我们不是把教材做成小游戏，而是让更多人拥有把自己的知识做成学习产品的能力。',.9,3.3,6.5,1.0,17,MUTED); pill(s,'拾知小镇 · Knowledge Town',.9,5.5,3.2,GREEN,WHITE,13); txt(s,'WaytoAGI 1024 · OPC',.9,6.25,3.4,.3,11,GREEN2,True); foot(s,15)
os.makedirs('docs',exist_ok=True); prs.save(OUT); print('saved',OUT,os.path.getsize(OUT))
